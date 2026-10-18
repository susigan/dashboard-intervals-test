"""Exportacao Excel diaria (TRIMP, MMP do dia, CP recalculado). Sinteticos, offline, sem producao.

    python -m pytest tests/pmc/test_export_cp_diario.py
"""
import hashlib
import json
import os
import sys
import tempfile
import unittest
from datetime import date
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
for _p in (HERE, ROOT, os.path.join(ROOT, 'utils'), os.path.join(ROOT, 'tools')):
    if _p not in sys.path:
        sys.path.insert(0, _p)

os.environ.setdefault('INTERVALS_ICU_API_KEY', 'ficticia-sem-rede')
import cp_model as cpm  # noqa: E402
import export_cp_diario as X  # noqa: E402
import fase1_util as U  # noqa: E402

MOD = 'Bike'                                    # classicas 300/720/1200; extra 60/180/300/720/1200
DUR = [60, 180, 300, 720, 1200, 3600]


def curva(dia, pontos, aid=None):
    """Linha de power_curves: pontos {duracao: W} (monotona decrescente com a duracao)."""
    secs = sorted(pontos)
    return {'date': dia, 'secs': secs, 'watts': [pontos[s] for s in secs], 'season': dia[:4],
            'activity_id': aid or 'a' + dia}


def pts(w60, w180, w300, w720, w1200, w3600):
    return dict(zip(DUR, (w60, w180, w300, w720, w1200, w3600)))


# 2025: season completa (inclui 3600 s: no modo coerente a season so e completa com as seis duracoes)
BASE_2025 = [curva('2025-06-01', pts(500, 400, 330, 290, 270, 240)), curva('2025-09-01', pts(480, 390, 320, 285, 262, 235))]
JAN10 = curva('2026-01-10', pts(450, 360, 300, 262, 245, 215))            # season 2026 completa -> teste elegivel
FEV01 = curva('2026-02-01', pts(300, 250, 230, 210, 200, 180))            # abaixo do melhor da season: nao e novo ponto
MAR01 = curva('2026-03-01', pts(470, 380, 320, 270, 250, 220))            # novo melhor em 300/720/1200
MAR10 = curva('2026-03-10', pts(520, 380, 300, 250, 230, 210))            # so 60 s sobe: extra, nao elegivel (modelo classico)
ABR01 = curva('2026-04-01', pts(480, 390, 330, 290, 270, 235))            # futuro: sobe 300/720/1200 de novo
REGISTOS = BASE_2025 + [JAN10, FEV01, MAR01, MAR10, ABR01]


def linhas(registos=REGISTOS, ate=date(2026, 4, 10), trimp=None, ausente=None, **kw):
    return X.linhas_modalidade(MOD, registos, trimp or {}, ausente or {}, [], date(2026, 1, 1), ate, **kw)


def por_data(ls):
    return {L['data'].isoformat(): L for L in ls}


class Calendario(unittest.TestCase):
    def test_1_todas_as_datas_sem_lacunas(self):
        ls = linhas()
        datas = [L['data'] for L in ls]
        self.assertEqual(datas[0], date(2026, 1, 1))
        self.assertEqual(datas[-1], date(2026, 4, 10))
        self.assertEqual(len(datas), (date(2026, 4, 10) - date(2026, 1, 1)).days + 1)
        self.assertEqual(len(set(datas)), len(datas))
        self.assertEqual(X.calendario(date(2026, 1, 1), date(2026, 1, 1)), [date(2026, 1, 1)])
        with self.assertRaises(ValueError):
            X.calendario(date(2026, 1, 2), date(2026, 1, 1))


class Trimp(unittest.TestCase):
    def test_2_dias_sem_treino_ficam_vazios_e_nunca_zero(self):
        ls = por_data(linhas(trimp={MOD: {'2026-01-10': 55.5}}))
        self.assertIsNone(ls['2026-01-09']['trimp'])
        self.assertEqual(ls['2026-01-10']['trimp'], 55.5)
        self.assertTrue(all(L['trimp'] != 0 for L in ls.values()))

    def test_agregacao_do_projeto_soma_por_dia_e_modalidade_e_ignora_fc_ignorada(self):
        ativ = [{'id': 'a', 'type': 'Ride', 'start_date_local': '2026-01-10T07:00:00', 'trimp': 40.0},
                {'id': 'b', 'type': 'VirtualRide', 'start_date_local': '2026-01-10T18:00:00', 'trimp': 15.5},
                {'id': 'c', 'type': 'Rowing', 'start_date_local': '2026-01-10T09:00:00', 'trimp': 99.0},
                {'id': 'd', 'type': 'Ride', 'start_date_local': '2026-01-11T09:00:00', 'trimp': 70.0, 'icu_ignore_hr': True},
                {'id': 'e', 'type': 'Ride', 'start_date_local': '2026-01-12T09:00:00'}]
        carga, ausente = X.trimp_por_dia(ativ)
        self.assertEqual(carga['Bike'], {'2026-01-10': 55.5})                   # nao mistura Row; soma o mesmo dia
        self.assertEqual(carga['Row'], {'2026-01-10': 99.0})
        self.assertNotIn('2026-01-11', carga['Bike'])                             # FC ignorada: ausente, nao zero
        self.assertEqual(ausente[('Bike', '2026-01-11')], ['fc_ignorada'])
        self.assertEqual(ausente[('Bike', '2026-01-12')], ['sem_valor'])
        ls = por_data(linhas(trimp=carga, ausente=ausente))
        self.assertIsNone(ls['2026-01-11']['trimp'])


class Testes(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ls = por_data(linhas())

    def test_3_dias_sem_teste_nao_mostram_cp_novo(self):
        for d in ('2026-01-01', '2026-01-09', '2026-01-15', '2026-02-02', '2026-03-05'):
            L = self.ls[d]
            self.assertEqual(L['estado'], 'sem_teste', d)
            self.assertIsNone(L['cp_w'], d)
            self.assertIsNone(L['texto'], d)
        # dia com atividade cujo MMP nao e um novo ponto: continua a ser sem_teste e sem CP
        L = self.ls['2026-02-01']
        self.assertEqual((L['estado'], L['cp_w']), ('sem_teste', None))
        self.assertTrue(any('nao e um novo ponto' in o for o in L['obs']))
        self.assertEqual(L['mmp_por_duracao'][300], 230.0)                        # o MMP do dia fica visivel

    def test_4_teste_elegivel_altera_o_conjunto_a_partir_da_data_correta(self):
        antes, dia, depois = self.ls['2026-03-01'], self.ls['2026-03-01'], self.ls['2026-03-02']
        self.assertEqual(dia['estado'], 'cp_calculado_com_aviso')                 # 3 pontos: avisos do CP Model
        self.assertIsNotNone(dia['cp_w'])
        self.assertIn('300 s: 320 W (2026-03-01)', dia['pontos'])
        self.assertIsNone(depois['cp_w'])                                         # nao se copia o CP para o dia seguinte
        jan = self.ls['2026-01-10']
        self.assertIn('300 s: 300 W (2026-01-10)', jan['pontos'])
        self.assertNotIn('2026-03-01', jan['pontos'])
        self.assertNotEqual(jan['cp_w'], dia['cp_w'])
        self.assertEqual(dia['dur_teste_s'], 1200)
        self.assertEqual(dia['mmp_dia_w'], 250.0)

    def test_5_teste_de_duracao_nao_elegivel_nao_entra_no_calculo(self):
        L = self.ls['2026-03-10']
        self.assertEqual(L['estado'], 'teste_nao_elegivel')
        self.assertIsNone(L['cp_w'])
        self.assertEqual(L['dur_teste_s'], 60)
        self.assertIn('nao entrou', L['texto'])
        self.assertTrue(any('3 parametros' in o for o in L['obs']))

    def test_6_resultado_futuro_nao_altera_o_passado(self):
        com = linhas(REGISTOS, ate=date(2026, 4, 10))
        sem = linhas([r for r in REGISTOS if r is not ABR01], ate=date(2026, 3, 31))
        for a, b in zip(com, sem):
            self.assertEqual(a, b, a['data'])
        self.assertEqual(por_data(com)['2026-04-01']['estado'], 'cp_calculado_com_aviso')   # e o futuro so o muda no seu dia

    def test_7_reproduz_o_cp_model_com_as_mesmas_entradas(self):
        chamadas = []

        def espia(dados, modalidade, min_pts=3):
            chamadas.append(json.dumps(dados.get('all_mmp_pts')))
            return cpm.calcular_cp_completo(dados, modalidade, min_pts=min_pts)
        ls = linhas(calc_fn=espia)
        calculados = [L for L in ls if L['cp_w'] is not None]
        self.assertEqual(len(chamadas), len(calculados))                          # so corre o CP em dias com teste
        for L in calculados:
            d = L['data'].isoformat()
            dados = cpm.pontos_de_curvas([r for r in REGISTOS if r['date'] <= d], MOD, season_activa=d[:4], modo='coerente')
            res = cpm.calcular_cp_completo(dados, MOD)
            self.assertEqual(L['cp_w'], res['melhor']['cp'], d)
            self.assertEqual(L['modelo'], res['melhor']['nome'], d)
            self.assertEqual(L['see_pct'], res['melhor']['see_pct'], d)
            self.assertEqual(L['wprime_j'], res['melhor']['wp'], d)
            self.assertEqual(L['cp_m2'], res['modelos']['M2 (WLS-1/t)']['cp'], d)

    def test_modelo_fixado_e_respeitado(self):
        ls = por_data(linhas(modelo='M2 (WLS-1/t)'))
        L = ls['2026-03-01']
        self.assertEqual(L['modelo'], 'M2 (WLS-1/t)')
        self.assertEqual(L['cp_w'], L['cp_m2'])

    def test_10_dados_insuficientes_e_erro_sao_estados_proprios_sem_cp(self):
        pouco = [curva('2026-01-10', {300: 300, 720: 262})]                       # so 2 duracoes classicas
        L = por_data(linhas(pouco, ate=date(2026, 1, 12)))['2026-01-10']
        self.assertEqual(L['estado'], 'dados_insuficientes')
        self.assertIsNone(L['cp_w'])
        self.assertIn('CP nao calculavel', L['texto'])

        def rebenta(*a, **k):
            raise RuntimeError('falha simulada')
        L = por_data(linhas(calc_fn=rebenta, ate=date(2026, 1, 12)))['2026-01-10']
        self.assertEqual(L['estado'], 'erro_calculo')
        self.assertIsNone(L['cp_w'])
        self.assertTrue(any('falha simulada' in o for o in L['obs']))
        self.assertIn(L['estado'], X.ESTADOS)


class ExcelEBanco(unittest.TestCase):
    def test_9_excel_legivel_com_abas_colunas_e_formatos(self):
        import openpyxl
        ls = {m: linhas() for m in X.MODALIDADES}
        with tempfile.TemporaryDirectory() as t:
            caminho = X.escrever_excel(ls, os.path.join(t, 'cp.xlsx'))
            wb = openpyxl.load_workbook(caminho)
        self.assertEqual(wb.sheetnames, ['Bike', 'Row', 'Run', 'Ski'])
        ws = wb['Bike']
        self.assertEqual(ws.max_row, 1 + len(ls['Bike']))
        cab = [c.value for c in ws[1]]
        self.assertEqual(cab[:6], ['Data', 'TRIMP', 'Melhor MMP / CP recalculado', 'Duração do teste do dia (s)',
                                   'MMP do dia (W)', 'CP recalculado (W)'])
        self.assertEqual(cab[6:12], ['Estado do cálculo', 'Modelo CP', 'Nº de pontos', 'W′ (J)', 'SEE (%)',
                                     'Observação / diagnóstico'])
        datas = [r[0].value for r in ws.iter_rows(min_row=2)]
        self.assertTrue(all(isinstance(d, __import__('datetime').datetime) for d in datas))
        self.assertEqual(ws['A2'].number_format, 'dd/mm/yyyy')
        linha_cp = next(r for r in ws.iter_rows(min_row=2) if r[6].value == 'cp_calculado_com_aviso')
        self.assertEqual(linha_cp[5].number_format, '0.0')
        self.assertIsInstance(linha_cp[5].value, float)
        self.assertIn('CP recalculado =', linha_cp[2].value)
        vazia = next(r for r in ws.iter_rows(min_row=2) if r[6].value == 'sem_teste')
        self.assertIsNone(vazia[1].value)                                          # nada de zeros
        self.assertIsNone(vazia[5].value)

    def test_8_nenhum_cp_e_gravado_e_a_base_original_fica_intacta(self):
        antes_fontes = open(X.__file__, encoding='utf-8').read()
        for proibido in ('INSERT', 'UPDATE ', 'DELETE FROM', 'api_perfil_guardar', 'upload(', 'requests.'):
            self.assertNotIn(proibido, antes_fontes, proibido)
        import db
        with U.base_temporaria(db) as caminho:
            db.upsert_power_curves([{'activity_id': r['activity_id'], 'type': MOD, 'date': r['date'], 'weight': None,
                                     'secs': r['secs'], 'watts': r['watts']} for r in REGISTOS])
            db.upsert_activities([{'id': 'x1', 'date': '2026-01-10', 'type': MOD, 'type_raw': 'Ride',
                                   'raw': {'id': 'x1', 'type': 'Ride', 'start_date_local': '2026-01-10T07:00:00',
                                           'trimp': 42.0}}])
            h0 = hashlib.sha256(open(caminho, 'rb').read()).hexdigest()
            n_cp = lambda: db._exec("SELECT name FROM sqlite_master WHERE name='cp_resultados'", fetch='all')  # noqa: E731
            self.assertEqual(n_cp(), [])
            db._conn.commit()
            h0 = hashlib.sha256(open(caminho, 'rb').read()).hexdigest()
            with tempfile.TemporaryDirectory() as t:
                ls = X.exportar(caminho, os.path.join(t, 'o.xlsx'), date(2026, 1, 1), date(2026, 1, 20))
                self.assertTrue(os.path.exists(os.path.join(t, 'o.xlsx')))
            self.assertEqual(hashlib.sha256(open(caminho, 'rb').read()).hexdigest(), h0)   # original intacta
            self.assertEqual(n_cp(), [])                                                   # nenhum cp_resultados criado
            self.assertEqual(ls['Bike'][9]['trimp'], 42.0)
            self.assertEqual(ls['Bike'][9]['estado'], 'cp_calculado_com_aviso')
            self.assertEqual(set(ls), set(X.MODALIDADES))
            self.assertTrue(all(len(v) == 20 for v in ls.values()))


if __name__ == '__main__':
    unittest.main()
