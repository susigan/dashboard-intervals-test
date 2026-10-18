"""Exportacao Excel DIARIA por modalidade: TRIMP, MMP do dia e CP recalculado (FERRAMENTA OFFLINE, so leitura).

    python tools/export_cp_diario.py --db copia_local.db --out cp_diario.xlsx [--desde 2026-01-01] [--ate AAAA-MM-DD]
                                     [--modelo auto|"M1 (WLS-P)"|...] [--modo coerente|season|recuo]
                                     [--marcos 2026-01-01=2026 --marcos ...]

Quatro abas (Bike, Row, Run, Ski), uma linha por dia do calendario, inclusive dias sem treino e sem teste. NAO reimplementa
nada do CP: usa `api_cp._registos` (as curvas por modalidade), `utils/cp_model.pontos_de_curvas` (selecao dos MMP e
duplicados), `utils/cp_model.calcular_cp_completo` (todos os modelos, SEE%, validacao), `config.season_de` e
`distribuicao_carga.agregar(metodo='trimp')`. Nao altera o CP Model, nao grava nada (nem cp_resultados) e nao acede a
Intervals.icu nem a producao: trabalha numa COPIA temporaria da base indicada.

Semantica (a mesma da gravacao do CP, `/api/perfil/guardar`: season da data, modo 'coerente', min_pts 3, duplicados
excluidos, sem override de pmax):
  * Para a data D so entram power_curves com data <= D (sem fuga do futuro) e a season e a de D.
  * "Teste do dia" = o dia produziu um MMP que o CP Model SELECIONA (o ponto escolhido para essa duracao tem a data D)
    ou o conjunto de pontos do modelo mudou nesse dia. Dia com atividade cujo MMP nao e um novo ponto: `sem_teste`.
  * O CP so e recalculado (e so aparece na coluna F) em dias com teste elegivel; nunca se copia o CP anterior.
Limites: a data de disponibilidade de um teste e a data da ATIVIDADE (nao ha data de deteccao/sincronizacao confiavel),
portanto a tabela mostra o que o algoritmo produziria com os dados que HOJE existem ate D, nao o que era conhecido em
tempo real. Os marcos de season (SEASON_START) vem da API e aqui nao se acede: sem `--marcos` usa-se SEASON_INICIO_MES.
"""
import argparse
import os
import shutil
import sys
import tempfile
from datetime import date, datetime, timedelta

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
for _p in (ROOT, os.path.join(ROOT, 'utils')):
    if _p not in sys.path:
        sys.path.insert(0, _p)

MODALIDADES = ('Bike', 'Row', 'Run', 'Ski')
DESDE_PADRAO = date(2026, 1, 1)
MIN_PTS = 3

ESTADOS = ('sem_teste', 'teste_nao_elegivel', 'dados_insuficientes', 'cp_calculado', 'cp_calculado_com_aviso',
           'erro_calculo')

# modelos cujo resultado depende do conjunto "completo" (3 parametros) e nao so do classico (calcular_cp_completo)
MODELOS_3P = ('3p hiperbolico', 'Ward-Smith', 'OM3CP', 'OMExp')


# ── calendario ────────────────────────────────────────────────────────────────────────────────────

def calendario(ini, fim):
    """Todos os dias de ini a fim, inclusive, sem lacunas."""
    if fim < ini:
        raise ValueError('fim anterior ao inicio')
    return [ini + timedelta(days=i) for i in range((fim - ini).days + 1)]


# ── dados (E/S isolada aqui; o resto e puro) ──────────────────────────────────────────────────────────

def carregar(db_path, marcos=None):
    """-> (registos_por_modalidade, atividades, marcos). Trabalha numa COPIA temporaria da base (a original nunca e
    aberta para escrita) e impede qualquer acesso a rede (os marcos de season vem de `marcos`, nunca da API)."""
    from unittest import mock
    tmp = tempfile.mkdtemp(prefix='cp_diario_')
    copia = os.path.join(tmp, os.path.basename(db_path))
    import sqlite3
    origem = sqlite3.connect(f'file:{os.path.abspath(db_path)}?mode=ro', uri=True)    # leitura; inclui o WAL
    destino = sqlite3.connect(copia)
    try:
        origem.backup(destino)
    finally:
        destino.close()
        origem.close()
    os.environ['SQLITE_PATH'] = copia
    os.environ.setdefault('INTERVALS_ICU_API_KEY', 'ficticia-sem-rede')
    import api_cp
    import db
    marcos = sorted(marcos or [])
    antes = (db._conn, db.DRIVER, db.ENABLED, db.SQLITE_PATH, db.DATABASE_URL)
    db.DATABASE_URL, db.SQLITE_PATH, db._conn = '', copia, None      # aponta o modulo db para a COPIA (e repoe no fim)
    try:
        db._connect()
        with mock.patch('api_client.seasons_do_atleta', return_value=marcos):
            registos = {m: api_cp._registos(m)[0] or [] for m in MODALIDADES}
        atividades = db.load_activities()
        try:
            db._conn.close()
        except Exception:
            pass
    finally:
        db._conn, db.DRIVER, db.ENABLED, db.SQLITE_PATH, db.DATABASE_URL = antes
        shutil.rmtree(tmp, ignore_errors=True)
    return registos, atividades, marcos


def trimp_por_dia(atividades):
    """-> ({mod: {AAAA-MM-DD: soma do TRIMP valido}}, {(mod, dia): [motivos de TRIMP ausente]}). Definicao do projeto:
    campo `trimp` da atividade (Intervals.icu), soma por dia e modalidade, FC ignorada/ausente/invalida nao conta
    (nunca vira zero)."""
    import distribuicao_carga as dc
    ag = dc.agregar(atividades, 'trimp')
    recs, _ = dc.classificar(atividades)
    ausentes = {}
    for r in recs:
        if r['trimp'] is None:
            ausentes.setdefault((r['mod'], r['dia']), []).append(r['trimp_motivo'])
    return ag['carga'], ausentes


# ── MMP e CP (so reutiliza o CP Model) ──────────────────────────────────────────────────────────────

def duracoes_da(modalidade):
    import cp_model as cpm
    return sorted(set(cpm.DURACOES_CLASSICAS[modalidade]) | set(cpm.DURACOES_EXTRA[modalidade]) | {3600})


def mmp_do_dia(registos_do_dia, duracoes):
    """Melhor MMP por duracao nas curvas do dia, com a mesma leitura (`_indice_curva`/`_watts_em`, tolerancia de 5 %)
    que o CP Model usa. Dia sem curva -> {}."""
    import perfil_metabolico as pmet
    melhor = {}
    for l in registos_do_dia:
        _d, _s, indice = pmet._indice_curva(l)
        if not indice:
            continue
        for d in duracoes:
            w = pmet._watts_em(indice, d)
            if w is not None and (d not in melhor or w > melhor[d]):
                melhor[d] = float(w)
    return melhor


def pontos_ate(registos, modalidade, dia, marcos, modo='coerente', pontos_fn=None):
    """Conjunto de pontos do CP Model para a data `dia`: so curvas com data <= dia, season da data (config.season_de)."""
    import cp_model as cpm
    from config import season_de
    pontos_fn = pontos_fn or cpm.pontos_de_curvas
    d = dia.isoformat()
    return pontos_fn([r for r in registos if r.get('date') and r['date'] <= d], modalidade,
                     season_activa=season_de(d, marcos or None), modo=modo)


def _assinatura(dados, modelo):
    """O que o modelo escolhido consome: classicos/2p -> all_mmp_pts; 3 parametros -> all_mmp_pts_full (+ pmax)."""
    if modelo in MODELOS_3P:
        return (tuple(dados.get('all_mmp_pts_full') or []), dados.get('pmax'))
    return (tuple(dados.get('all_mmp_pts') or []),)


def _resumo_pontos(dados):
    datas = dados.get('datas') or {}
    return ' | '.join(f"{int(t)} s: {w:.0f} W ({datas.get(str(int(t))) or '?'})" for w, t in (dados.get('all_mmp_pts') or []))


def _escolher_modelo(res, modelo):
    """(nome, dict do modelo) segundo o pedido: 'auto' = o `melhor` (menor SEE% entre os classicos), como a gravacao
    do CP Model quando o utilizador nao fixa modelo."""
    if modelo in (None, 'auto'):
        m = res.get('melhor') or {}
        return m.get('nome'), (res.get('modelos') or {}).get(m.get('nome')) or {}
    return modelo, (res.get('modelos') or {}).get(modelo) or {}


def linhas_modalidade(modalidade, registos, trimp, trimp_ausente, marcos, ini, fim, modelo='auto', modo='coerente',
                      min_pts=MIN_PTS, calc_fn=None, pontos_fn=None):
    """Uma linha por dia de ini a fim. Puro: recebe registos/TRIMP ja carregados. `calc_fn`/`pontos_fn` so existem para
    os testes provarem que SAO as funcoes do CP Model que correm (por omissao: `calcular_cp_completo`/`pontos_de_curvas`)."""
    import cp_model as cpm
    from config import season_de
    calc_fn = calc_fn or cpm.calcular_cp_completo
    classicas = set(cpm.DURACOES_CLASSICAS[modalidade])
    extras = set(cpm.DURACOES_EXTRA[modalidade])
    durs = duracoes_da(modalidade)
    por_dia = {}
    for r in registos:
        if r.get('date'):
            por_dia.setdefault(r['date'], []).append(r)

    prev = pontos_ate(registos, modalidade, ini - timedelta(days=1), marcos, modo, pontos_fn)
    prev_season = season_de((ini - timedelta(days=1)).isoformat(), marcos or None)
    linhas = []
    for dia in calendario(ini, fim):
        iso = dia.isoformat()
        sea = season_de(iso, marcos or None)
        rows = por_dia.get(iso, [])
        mmp_dia = mmp_do_dia(rows, durs)
        L = {'data': dia, 'trimp': (trimp.get(modalidade) or {}).get(iso), 'texto': None, 'dur_teste_s': None,
             'mmp_dia_w': None, 'cp_w': None, 'estado': 'sem_teste', 'modelo': None, 'n_pontos': None, 'wprime_j': None,
             'see_pct': None, 'obs': [], 'cp_m1': None, 'cp_m2': None, 'cp_m3': None, 'season_conjunto': None,
             'pontos': None, 'mmp_por_duracao': {d: mmp_dia.get(d) for d in durs}}
        mot = (trimp_ausente or {}).get((modalidade, iso))
        if L['trimp'] is None and mot:
            L['obs'].append('atividade sem TRIMP valido (' + ', '.join(sorted(set(mot))) + ')')
        if not rows and sea == prev_season:
            linhas.append(L)
            continue
        atual = pontos_ate(registos, modalidade, dia, marcos, modo, pontos_fn)
        nome_modelo = modelo if modelo not in (None, 'auto') else 'M1 (WLS-P)'      # so decide classico vs 3 parametros
        mudou = _assinatura(atual, nome_modelo) != _assinatura(prev, nome_modelo)
        datas = atual.get('datas') or {}
        entrou = sorted(d for d in durs if rows and d in mmp_dia and datas.get(str(d)) == iso)
        prev, prev_season = atual, sea
        if not rows:
            if mudou:
                L['obs'].append('conjunto de pontos mudou por troca de season (sem teste)')
            linhas.append(L)
            continue
        if not entrou and not mudou:
            L['obs'].append('MMP do dia nao e um novo ponto do CP (abaixo do melhor ja selecionado)')
            linhas.append(L)
            continue
        # ha um resultado do dia que o CP Model selecionou, ou o conjunto mudou neste dia
        usados = {int(t) for _w, t in (atual.get('all_mmp_pts') or [])} | \
                 ({int(t) for _w, t in (atual.get('all_mmp_pts_full') or [])} if nome_modelo in MODELOS_3P else set())
        elegiveis = [d for d in entrou if d in usados]
        if entrou:
            principal = max(elegiveis or entrou)
            L['dur_teste_s'], L['mmp_dia_w'] = principal, mmp_dia.get(principal)
        partes = [f'{d} s = {mmp_dia[d]:.0f} W' + ('' if d in usados else ' (fora do ajuste)') for d in entrou]
        if not mudou:
            motivos = []
            for d in entrou:
                if d == 3600:
                    motivos.append('3600 s nunca entra no ajuste (so validacao)')
                elif d in extras and d not in classicas and nome_modelo not in MODELOS_3P:
                    motivos.append(f'{d} s so alimenta modelos de 3 parametros (informativos)')
                else:
                    motivos.append(f'{d} s descartado pela regra de duplicados')
            L['estado'] = 'teste_nao_elegivel'
            L['obs'] += ['nao entrou no calculo: ' + '; '.join(motivos)]
            L['texto'] = 'MMP ' + ' + '.join(partes) + '; nao entrou no CP'
            linhas.append(L)
            continue
        L['season_conjunto'], L['pontos'] = atual.get('season_do_conjunto'), _resumo_pontos(atual)
        if not entrou:
            L['obs'].append('o conjunto de pontos mudou neste dia sem um novo melhor (regra de conjunto da season)')
        elif not elegiveis:
            L['obs'].append('o resultado do dia nao esta no conjunto do modelo, mas o conjunto de pontos mudou neste dia '
                            '(no modo coerente a season so e "completa" quando tem tambem o MMP de 3600 s)')
        try:
            res = calc_fn(atual, modalidade, min_pts=min_pts)
        except Exception as e:                                           # nao inventar CP: erro e estado proprio
            L['estado'] = 'erro_calculo'
            L['obs'].append(f'{type(e).__name__}: {e}')
            L['texto'] = ('MMP ' + ' + '.join(partes) + '; ' if partes else '') + 'erro no calculo do CP'
            linhas.append(L)
            continue
        for k, nm in (('cp_m1', 'M1 (WLS-P)'), ('cp_m2', 'M2 (WLS-1/t)'), ('cp_m3', 'M3 (NL-2p)')):
            L[k] = ((res.get('modelos') or {}).get(nm) or {}).get('cp')
        nome, m = _escolher_modelo(res, modelo)
        if not res.get('ok') and not (res.get('modelos')):
            L['estado'] = 'dados_insuficientes'
            L['obs'].append(res.get('motivo') or 'dados insuficientes')
            L['n_pontos'] = res.get('n_mmp')
            L['texto'] = ('MMP ' + ' + '.join(partes) + '; ' if partes else '') + 'CP nao calculavel (dados insuficientes)'
        elif not m.get('cp'):
            L['estado'] = 'erro_calculo'
            L['modelo'] = nome
            L['obs'].append(f'o modelo {nome!r} nao produziu CP com estes pontos')
            L['texto'] = ('MMP ' + ' + '.join(partes) + '; ' if partes else '') + 'modelo sem resultado'
        else:
            avisos = [a for a in (res.get('validacao') or []) if a.get('gravidade') in ('alto', 'medio')]
            L.update(cp_w=m['cp'], modelo=nome, n_pontos=m.get('n_pts'), wprime_j=m.get('wp'), see_pct=m.get('see_pct'),
                     estado='cp_calculado_com_aviso' if avisos else 'cp_calculado')
            L['obs'] += [f"{a.get('gravidade')}: {a.get('chave')}" for a in avisos]
            L['texto'] = ('MMP ' + ' + '.join(partes) + '; ' if partes else 'conjunto mudou; ') + \
                f'CP recalculado = {m["cp"]:.1f} W'
        linhas.append(L)
    return linhas


# ── Excel ─────────────────────────────────────────────────────────────────────────────────────────

CABECALHO = ['Data', 'TRIMP', 'Melhor MMP / CP recalculado', 'Duração do teste do dia (s)', 'MMP do dia (W)',
             'CP recalculado (W)', 'Estado do cálculo', 'Modelo CP', 'Nº de pontos', 'W′ (J)', 'SEE (%)',
             'Observação / diagnóstico', 'CP M1 (W)', 'CP M2 (W)', 'CP M3 (W)', 'Season do conjunto',
             'Pontos usados (data do ponto)']

NOTAS = {
    'B': 'TRIMP do dia na modalidade (campo trimp da atividade, Intervals.icu; soma por dia). Vazio = sem atividade ou sem TRIMP valido (nunca zero).',
    'C': 'Texto legivel: MMP do teste do dia e CP recalculado com os dados ate esta data.',
    'F': 'CP recalculado pelo CP Model com os pontos disponiveis ate esta data. Vazio = sem novo calculo (nao se copia o CP anterior).',
    'G': 'sem_teste | teste_nao_elegivel | dados_insuficientes | cp_calculado | cp_calculado_com_aviso | erro_calculo',
    'H': 'Modelo usado. auto = o de menor SEE% entre M1/M2/M3 (como a gravacao do CP Model sem modelo fixado).',
    'L': 'Avisos da validacao do CP Model e motivos de nao entrada. Recalculo retrospetivo: NAO e o CP confirmado nem o P0 do IR.',
}


def escrever_excel(linhas_por_mod, caminho, meta=None):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    cores = {'cp_calculado': 'E8F5E9', 'cp_calculado_com_aviso': 'FFF8E1', 'teste_nao_elegivel': 'ECEFF1',
             'dados_insuficientes': 'FFEBEE', 'erro_calculo': 'FFCDD2'}
    for mod in MODALIDADES:
        linhas = linhas_por_mod[mod]
        durs = sorted(linhas[0]['mmp_por_duracao']) if linhas else []
        ws = wb.create_sheet(mod)
        cab = CABECALHO + [f'MMP do dia {d} s (W)' for d in durs]
        ws.append(cab)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical='center')
        for col, nota in NOTAS.items():
            ws[f'{col}1'].comment = Comment(nota, 'export_cp_diario')
        for L in linhas:
            ws.append([L['data'], L['trimp'], L['texto'], L['dur_teste_s'], L['mmp_dia_w'], L['cp_w'], L['estado'],
                       L['modelo'], L['n_pontos'], L['wprime_j'], L['see_pct'], '; '.join(L['obs']) or None,
                       L['cp_m1'], L['cp_m2'], L['cp_m3'], L['season_conjunto'], L['pontos']]
                      + [L['mmp_por_duracao'].get(d) for d in durs])
            r = ws.max_row
            ws.cell(r, 1).number_format = 'dd/mm/yyyy'
            for col, fmt in ((2, '0.0'), (5, '0.0'), (6, '0.0'), (10, '#,##0'), (11, '0.00'), (13, '0.0'), (14, '0.0'),
                             (15, '0.0')):
                ws.cell(r, col).number_format = fmt
            for k in range(len(durs)):
                ws.cell(r, len(CABECALHO) + 1 + k).number_format = '0.0'
            if L['estado'] in cores:
                for col in range(1, len(CABECALHO) + 1):
                    ws.cell(r, col).fill = PatternFill('solid', fgColor=cores[L['estado']])
        larguras = [12, 9, 46, 12, 11, 12, 22, 16, 9, 10, 9, 60, 10, 10, 10, 14, 52]
        for i, w in enumerate(larguras + [12] * len(durs), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 45
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(cab))}{ws.max_row}'
    if meta:
        wb.properties.description = '; '.join(f'{k}={v}' for k, v in meta.items())
    wb.save(caminho)
    return caminho


# ── linha de comando ──────────────────────────────────────────────────────────────────────────────

def _parse_marcos(itens):
    out = []
    for x in itens or []:
        d, _, nome = x.partition('=')
        out.append((d, nome or d[:4]))
    return out


def exportar(db_path, saida, desde=DESDE_PADRAO, ate=None, modelo='auto', modo='coerente', marcos=None):
    ate = ate or date.today()
    registos, atividades, marcos = carregar(db_path, marcos)
    carga, ausente = trimp_por_dia(atividades)
    linhas = {m: linhas_modalidade(m, registos[m], carga, ausente, marcos, desde, ate, modelo, modo) for m in MODALIDADES}
    escrever_excel(linhas, saida, {'desde': desde, 'ate': ate, 'modelo': modelo, 'modo': modo,
                                   'gerado_em': datetime.now().isoformat(timespec='seconds')})
    return linhas


def _main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--db', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--desde', default=DESDE_PADRAO.isoformat())
    ap.add_argument('--ate', default=None)
    ap.add_argument('--modelo', default='auto')
    ap.add_argument('--modo', default='coerente', choices=['coerente', 'season', 'recuo'])
    ap.add_argument('--marcos', action='append', help='AAAA-MM-DD=nome da season (SEASON_START); repetir')
    a = ap.parse_args(argv)
    linhas = exportar(a.db, a.out, date.fromisoformat(a.desde), date.fromisoformat(a.ate) if a.ate else None,
                      a.modelo, a.modo, _parse_marcos(a.marcos))
    for m, ls in linhas.items():
        cont = {}
        for L in ls:
            cont[L['estado']] = cont.get(L['estado'], 0) + 1
        print(m, len(ls), 'linhas', dict(sorted(cont.items())))
    print('escrito', a.out)


if __name__ == '__main__':
    _main()
